#!/usr/bin/env python3
from __future__ import annotations
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "17365D"
MID_BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CODE_FILL = "F6F8FA"
FONT = "Sarabun"
MONO = "Liberation Mono"


def set_run_font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell_or_p, fill):
    target = cell_or_p._tc if hasattr(cell_or_p, "_tc") else cell_or_p._p
    pr_name = "w:tcPr" if hasattr(cell_or_p, "_tc") else "w:pPr"
    pr = target.find(qn(pr_name))
    if pr is None:
        pr = OxmlElement(pr_name)
        target.insert(0, pr)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("หน้า ")
    set_run_font(run, size=9, color="666666")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_inline(paragraph, text, size=10.5):
    # Lightweight Markdown inline rendering for bold, code and links.
    pat = re.compile(r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
    pos = 0
    for m in pat.finditer(text):
        if m.start() > pos:
            set_run_font(paragraph.add_run(text[pos:m.start()]), size=size)
        token = m.group(0)
        if token.startswith("**"):
            set_run_font(paragraph.add_run(token[2:-2]), size=size, bold=True)
        elif token.startswith("`"):
            set_run_font(paragraph.add_run(token[1:-1]), name=MONO, size=9.2, color=BLUE)
        else:
            label, url = re.match(r"\[([^]]+)\]\(([^)]+)\)", token).groups()
            set_run_font(paragraph.add_run(f"{label} ({url})"), size=size, color=MID_BLUE)
        pos = m.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), size=size)


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tcw = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tcw.set(qn("w:w"), str(widths[idx]))
            tcw.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    # Repeat the table header after an automatic page break.
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def create_architecture_png(path: Path):
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
    fig, ax = plt.subplots(figsize=(9.0, 4.4), dpi=180)
    ax.set_xlim(0, 10); ax.set_ylim(0, 5); ax.axis("off")
    boxes = {
        "IHP Pad Ring": (0.4, 1.6, 1.7, 1.7),
        "Raven Digital Core": (2.7, 1.6, 2.0, 1.7),
        "PicoRV32\nRV32IMC": (5.3, 3.35, 1.8, 1.0),
        "IHP SRAM\n1024 x 32": (7.7, 3.35, 1.8, 1.0),
        "UART / GPIO\nControl SPI": (5.3, 0.65, 1.8, 1.0),
        "SPI/QSPI\nFlash XIP": (7.7, 0.65, 1.8, 1.0),
    }
    for label, (x,y,w,h) in boxes.items():
        patch = FancyBboxPatch((x,y),w,h,boxstyle="round,pad=0.08",fc="#E8EEF5",ec="#17365D",lw=1.5)
        ax.add_patch(patch); ax.text(x+w/2,y+h/2,label,ha="center",va="center",fontsize=10,color="#17365D",weight="bold")
    edges = [((2.1,2.45),(2.7,2.45)),((4.7,2.45),(5.3,3.85)),((7.1,3.85),(7.7,3.85)),
             ((4.7,2.45),(5.3,1.15)),((7.1,1.15),(7.7,1.15))]
    for a,b in edges:
        ax.add_patch(FancyArrowPatch(a,b,arrowstyle="-|>",mutation_scale=12,color="#2E74B5",lw=1.5))
    ax.text(5,4.75,"Raven PicoRV32 — IHP SG13G2 Digital Full-Chip Baseline",ha="center",fontsize=12,color="#17365D",weight="bold")
    fig.tight_layout(); fig.savefig(path,bbox_inches="tight",facecolor="white"); plt.close(fig)


def main(src: Path, out: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(1.0)
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.header_distance = sec.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT; normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, before, after, color in (
        ("Title", 28, 0, 8, BLUE), ("Subtitle", 14, 0, 18, "555555"),
        ("Heading 1", 16, 18, 10, MID_BLUE), ("Heading 2", 13, 14, 7, MID_BLUE),
        ("Heading 3", 12, 10, 5, BLUE)):
        s = styles[name]; s.font.name = FONT; s.font.size = Pt(size); s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        s = styles[name]; s.font.name = FONT; s.font.size = Pt(10.5)
        s.paragraph_format.left_indent = Inches(0.375)
        s.paragraph_format.first_line_indent = Inches(-0.188)
        s.paragraph_format.space_after = Pt(4); s.paragraph_format.line_spacing = 1.25

    header = sec.header.paragraphs[0]
    set_run_font(header.add_run("RAVEN SOC | FULL-CHIP IMPLEMENTATION GUIDE"), size=8.5, bold=True, color="777777")
    add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(120); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("RAVEN"), size=34, bold=True, color=BLUE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("PicoRV32 Full-Chip Implementation"), size=22, bold=True, color=MID_BLUE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("LibreLane + IHP Open PDK ihp-sg13g2"), size=15, color="555555")
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(30); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("คู่มือเชิงปฏิบัติและชุดโค้ดพร้อมรัน\nDigital full-chip baseline • 100 MHz target • 4 KiB IHP SRAM"), size=12, color=BLUE)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(120); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("Revision 1.0 | 24 August 2026"), size=10, color="777777")
    doc.add_page_break()

    arch = out.parent / "raven_architecture.png"
    create_architecture_png(arch)

    lines = src.read_text(encoding="utf-8").splitlines()
    # Skip Markdown title/subtitle already represented on the cover.
    idx = 2
    in_code = False; code_lang = ""; code_lines = []
    while idx < len(lines):
        line = lines[idx]
        if line.startswith("```"):
            if not in_code:
                in_code = True; code_lang = line[3:].strip(); code_lines = []
            else:
                if code_lang == "mermaid":
                    picture = doc.add_picture(str(arch), width=Inches(6.25))
                    picture._inline.docPr.set("descr", "Raven PicoRV32 IHP SG13G2 digital full-chip architecture")
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0.15); p.paragraph_format.right_indent = Inches(0.15)
                    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(7)
                    p.paragraph_format.line_spacing = 1.0; shade(p, CODE_FILL)
                    set_run_font(p.add_run("\n".join(code_lines)), name=MONO, size=8.2, color="222222")
                in_code = False
            idx += 1; continue
        if in_code:
            code_lines.append(line); idx += 1; continue
        if line.startswith("| ") and idx + 1 < len(lines) and re.match(r"^\|[-: |]+\|$", lines[idx+1]):
            rows=[]
            while idx < len(lines) and lines[idx].startswith("|"):
                rows.append([x.strip() for x in lines[idx].strip().strip("|").split("|")]); idx += 1
            header_row=rows[0]; body=rows[2:]; cols=len(header_row)
            table=doc.add_table(rows=1, cols=cols); table.style="Table Grid"
            data=[header_row]+body
            for ridx,row in enumerate(data):
                cells=table.rows[0].cells if ridx==0 else table.add_row().cells
                for cidx,text in enumerate(row):
                    cells[cidx].text=""
                    add_inline(cells[cidx].paragraphs[0],text,size=9.2)
                    if ridx==0:
                        shade(cells[cidx],LIGHT_BLUE)
                        for run in cells[cidx].paragraphs[0].runs: run.bold=True
            # Wider narrative columns; exact total is 9360 DXA.
            if cols==2: widths=[2700,6660]
            elif cols==3: widths=[1700,2500,5160]
            elif cols==4: widths=[1200,2200,2600,3360]
            else: widths=[9360//cols]*cols; widths[-1]+=9360-sum(widths)
            set_table_geometry(table,widths)
            doc.add_paragraph().paragraph_format.space_after=Pt(2)
            continue
        if not line.strip():
            idx += 1; continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2"); idx += 1; continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 1"); idx += 1; continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:], style="Heading 1"); idx += 1; continue
        if line.startswith("- [ ] "):
            p=doc.add_paragraph(style="List Bullet"); add_inline(p,"☐ "+line[6:]); idx += 1; continue
        if line.startswith("- "):
            p=doc.add_paragraph(style="List Bullet"); add_inline(p,line[2:]); idx += 1; continue
        if re.match(r"^\d+\. ", line):
            p=doc.add_paragraph(style="List Number"); add_inline(p,re.sub(r"^\d+\. ","",line)); idx += 1; continue
        if line.startswith("> "):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.25); p.paragraph_format.right_indent=Inches(0.2)
            shade(p,LIGHT_GRAY); add_inline(p,line[2:]); idx += 1; continue
        # Join wrapped Markdown prose until another structural line.
        parts=[line.strip()]; idx += 1
        while idx < len(lines) and lines[idx].strip() and not re.match(r"^(#{1,3} |```|\| |- |\d+\. |> )",lines[idx]):
            parts.append(lines[idx].strip()); idx += 1
        p=doc.add_paragraph(); add_inline(p," ".join(parts))

    out.parent.mkdir(parents=True,exist_ok=True)
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
